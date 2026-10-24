"""Leitura e verificações automáticas de folhas de cálculo Excel (.xlsx)."""

from __future__ import annotations

import re

from openpyxl import load_workbook

from corretor.modelos import Verificacao

MAX_LINHAS, MAX_COLUNAS = 200, 30  # o que se mostra à IA por folha
ERROS = ("#DIV/0!", "#REF!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!")
# Rótulos que indicam uma célula calculada: o valor ao lado (ou por baixo) devia ser uma fórmula.
ROTULO_CALCULO = re.compile(r"\b(total|subtotal|soma|m[ée]dia|m[áa]ximo|m[íi]nimo|iva|lucro|saldo)\b", re.I)
# O .xlsx guarda as funções em inglês. Nomes no Excel em português, para mostrar ao professor.
FUNCOES_PT = {
    "SUM": "SOMA", "AVERAGE": "MÉDIA", "IF": "SE", "VLOOKUP": "PROCV", "HLOOKUP": "PROCH", "XLOOKUP": "PROCX",
    "COUNT": "CONTAR", "COUNTA": "CONTAR.VAL", "COUNTIF": "CONTAR.SE", "COUNTIFS": "CONTAR.SE.S",
    "SUMIF": "SOMA.SE", "SUMIFS": "SOMA.SE.S", "AVERAGEIF": "MÉDIA.SE", "MAX": "MÁXIMO", "MIN": "MÍNIMO",
    "ROUND": "ARRED", "AND": "E", "OR": "OU", "NOT": "NÃO", "TODAY": "HOJE", "NOW": "AGORA",
    "CONCATENATE": "CONCATENAR", "CONCAT": "CONCAT", "IFERROR": "SE.ERRO", "RANK": "ORDEM", "LEFT": "ESQUERDA",
    "RIGHT": "DIREITA", "LEN": "NÚM.CARAT", "UPPER": "MAIÚSCULAS", "INDEX": "ÍNDICE", "MATCH": "CORRESP",
}


def _funcoes(formula: str) -> list[str]:
    return [f.removeprefix("_xlfn.") for f in re.findall(r"([A-Z_][A-Z0-9_.]*)\(", formula.upper())]


def _titulo_grafico(grafico) -> str:
    try:
        return "".join(r.t for p in grafico.title.tx.rich.p for r in (p.r or []))
    except AttributeError:
        return ""


def _abrir(caminho: str):
    if caminho.lower().endswith(".xls"):
        raise ValueError("ficheiros .xls antigos não são suportados. Guarda como .xlsx no Excel.")
    # Duas leituras: uma com as fórmulas, outra com os valores que o Excel calculou da última vez.
    return load_workbook(caminho, data_only=False), load_workbook(caminho, data_only=True)


def ler_excel(caminho: str) -> str:
    """Texto com o conteúdo de cada folha, célula a célula, para a IA ler."""
    wb_f, wb_v = _abrir(caminho)
    partes = []
    for ws in wb_f.worksheets:
        wv = wb_v[ws.title]
        linhas = [f'<folha nome="{ws.title}" dimensao="{ws.dimensions}">']
        for row in ws.iter_rows(max_row=min(ws.max_row, MAX_LINHAS), max_col=min(ws.max_column, MAX_COLUNAS)):
            for c in row:
                if c.value is None:
                    continue
                valor = wv[c.coordinate].value
                if c.data_type == "f":
                    linhas.append(f"{c.coordinate}: ={str(c.value).lstrip('=')} -> {valor}")
                else:
                    formato = f" [formato {c.number_format}]" if c.number_format != "General" else ""
                    linhas.append(f"{c.coordinate}: {c.value!r}{formato}")
        if ws.max_row > MAX_LINHAS or ws.max_column > MAX_COLUNAS:
            linhas.append(f"(folha cortada: mostradas só as primeiras {MAX_LINHAS} linhas e {MAX_COLUNAS} colunas)")
        for g in ws._charts:
            linhas.append(f"Gráfico: {type(g).__name__}, título '{_titulo_grafico(g) or 'sem título'}'")
        for intervalo in ws.conditional_formatting:
            linhas.append(f"Formatação condicional em {intervalo.sqref}")
        if ws.merged_cells.ranges:
            linhas.append(f"Células unidas: {', '.join(str(r) for r in ws.merged_cells.ranges)}")
        linhas.append("</folha>")
        partes.append("\n".join(linhas))
    return "\n\n".join(partes)


def verificar_excel(caminho: str) -> list[Verificacao]:
    wb_f, wb_v = _abrir(caminho)
    res: list[Verificacao] = []
    n_formulas, funcoes, absolutas = 0, set(), 0
    erros, fixos = [], []
    n_graficos = n_condicional = 0

    for ws in wb_f.worksheets:
        wv = wb_v[ws.title]
        n_graficos += len(ws._charts)
        n_condicional += len(list(ws.conditional_formatting))
        for row in ws.iter_rows():
            for c in row:
                if c.data_type == "f":
                    n_formulas += 1
                    funcoes.update(_funcoes(str(c.value)))
                    absolutas += "$" in str(c.value)
                    v = wv[c.coordinate].value
                    if isinstance(v, str) and v in ERROS:
                        erros.append(f"{ws.title}!{c.coordinate} ({v})")
                elif isinstance(c.value, str) and ROTULO_CALCULO.search(c.value):
                    # Rótulo numa linha: confirma as células à direita.
                    for col in range(c.column + 1, min(c.column + 12, ws.max_column + 1)):
                        vizinha = ws.cell(c.row, col)
                        if isinstance(vizinha.value, (int, float)) and not isinstance(vizinha.value, bool):
                            fixos.append(f"{ws.title}!{vizinha.coordinate} ('{c.value.strip()}')")
                    # Rótulo no cabeçalho de uma coluna: confirma as células por baixo.
                    if c.row <= 3:
                        for r in range(c.row + 1, min(ws.max_row, c.row + 500) + 1):
                            abaixo = ws.cell(r, c.column)
                            if isinstance(abaixo.value, (int, float)) and not isinstance(abaixo.value, bool):
                                fixos.append(f"{ws.title}!{abaixo.coordinate} (coluna '{c.value.strip()}')")

    fixos = list(dict.fromkeys(fixos))
    res.append(Verificacao(n_formulas > 0, f"{n_formulas} fórmula(s) no livro." if n_formulas else "Não tem nenhuma fórmula."))
    if funcoes:
        nomes = sorted(FUNCOES_PT.get(f, f) for f in funcoes)
        res.append(Verificacao(True, f"Funções usadas: {', '.join(nomes)}."))
    if absolutas:
        res.append(Verificacao(True, f"{absolutas} fórmula(s) com referências absolutas ($)."))
    if fixos:
        lista = ", ".join(fixos[:8]) + (f" e mais {len(fixos) - 8}" if len(fixos) > 8 else "")
        res.append(Verificacao(False, f"Valores escritos à mão onde se espera uma fórmula: {lista}."))
    if erros:
        res.append(Verificacao(False, f"Células com erro: {', '.join(erros[:8])}."))
    res.append(Verificacao(n_graficos > 0, f"{n_graficos} gráfico(s)." if n_graficos else "Não tem gráficos."))
    if n_condicional:
        res.append(Verificacao(True, f"{n_condicional} regra(s) de formatação condicional."))
    folhas = [ws.title for ws in wb_f.worksheets]
    res.append(Verificacao(True, f"{len(folhas)} folha(s): {', '.join(folhas)}."))
    return res
